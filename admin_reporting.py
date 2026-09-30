"""Read-only dashboard and XLSX export; no payment or database mutations."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from flask import request, session, redirect, url_for, render_template, send_file, abort


def month_shift(value, offset):
    n = value.year * 12 + value.month - 1 + offset
    return value.replace(year=n // 12, month=n % 12 + 1, day=1)


def period_from_args(args, now=None):
    now = now or datetime.now()
    raw_month = args.get('month', '') or now.strftime('%Y-%m')
    try:
        end_month = datetime.strptime(raw_month, '%Y-%m')
        if not 1900 <= end_month.year <= 9998:
            raise ValueError()
        months = args.get('months', '1')
        if months != 'all' and months not in {str(i) for i in range(1, 13)}:
            raise ValueError()
    except ValueError:
        abort(400, 'Choose a valid month and a period from 1 to 12 months or All time.')
    if months == 'all':
        return dict(months=months, month=raw_month, start=None, end=None, label='All time')
    start = month_shift(end_month, 1 - int(months))
    end = month_shift(end_month, 1)
    label = start.strftime('%b %Y')
    if months != '1':
        label += ' – ' + end_month.strftime('%b %Y')
    return dict(months=months, month=raw_month, start=start.strftime('%Y-%m-%d'), end=end.strftime('%Y-%m-%d'), label=label)


def reporting_data(connection, period):
    where = ''
    params = []
    date_expr = "COALESCE(NULLIF(o.paid_at, ''), o.created_at)"
    if period['start']:
        where = f'WHERE datetime({date_expr}) >= datetime(?) AND datetime({date_expr}) < datetime(?)'
        params = [period['start'], period['end']]
    orders = [dict(row) for row in connection.execute(f'''
        SELECT o.*, c.email, c.first_name, c.last_name, c.phone,
               {date_expr} AS reporting_date,
               EXISTS(SELECT 1 FROM reports r WHERE r.order_id=o.id) AS has_report
        FROM orders o LEFT JOIN customers c ON c.id=o.customer_id
        {where} ORDER BY datetime({date_expr}) DESC, o.id DESC
    ''', params).fetchall()]
    customer_ids = {o['customer_id'] for o in orders if o['customer_id'] is not None}
    customers = [dict(row) for row in connection.execute('SELECT * FROM customers ORDER BY created_at DESC').fetchall() if not period['start'] or row['id'] in customer_ids]
    for customer in customers:
        customer['order_count'] = sum(o['customer_id'] == customer['id'] for o in orders)
    paid = [o for o in orders if o['status'] in ('paid', 'completed')]
    revenue_by_currency = {}
    for order in paid:
        currency = order.get('currency') or 'USD'
        revenue_by_currency[currency] = revenue_by_currency.get(currency, Decimal('0')) + Decimal(str(order['amount'] or 0))
    revenue = revenue_by_currency.get('USD', Decimal('0'))
    buckets = {}
    if period['start']:
        dt = datetime.strptime(period['start'], '%Y-%m-%d')
        while dt.strftime('%Y-%m-%d') < period['end']:
            buckets[dt.strftime('%Y-%m')] = {'orders':0,'sales':0,'revenue':Decimal('0')}
            dt = month_shift(dt, 1)
    for order in orders:
        key = ((order['reporting_date'] or '')[:7] or 'Unknown') + ' ' + (order.get('currency') or 'USD')
        bucket = buckets.setdefault(key, {'orders':0,'sales':0,'revenue':Decimal('0')})
        bucket['orders'] += 1
        if order['status'] in ('paid', 'completed'):
            bucket['sales'] += 1
            bucket['revenue'] += Decimal(str(order['amount'] or 0))
    monthly = [dict(month=k, currency=k.split()[-1] if ' ' in k else 'USD', **v) for k,v in sorted(buckets.items()) if v['orders']]
    max_revenue = max((m['revenue'] for m in monthly), default=Decimal('0'))
    for m in monthly:
        m['percent'] = float(m['revenue'] / max_revenue * 100) if max_revenue else 0
    return dict(period=period, orders=orders, customers=customers, monthly=monthly,
                total_orders=len(orders), total_revenue=revenue, revenue_by_currency=revenue_by_currency, paid_orders=len(paid),
                pending_reports=sum(o['status'] in ('paid','completed') and o.get('report_status','awaiting')!='delivered' for o in orders),
                total_customers=len(customers))


def build_workbook(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    rows = [ ['VINLookupNow Sales Report', data['period']['label']],
             ['Date basis', 'Payment date; order date if no payment date'],
             ['Revenue basis', 'Paid and completed orders; gross recorded revenue'],
             ['Total orders', data['total_orders']], ['Paid / completed orders', data['paid_orders']],
             ['Gross revenue (USD)', float(data['total_revenue'])],
             ['Pending reports', data['pending_reports']], ['Customers', data['total_customers']] ]
    for row in rows: summary.append(row)
    for currency, amount in data['revenue_by_currency'].items():
        if currency != 'USD': summary.append(['Gross revenue (' + currency + ')', float(amount)])
    summary['B6'].number_format = '#,##0.00'
    ws = wb.create_sheet('Monthly Sales')
    ws.append(['Month / Currency', 'All orders', 'Paid / completed', 'Gross revenue'])
    for m in data['monthly']: ws.append([m['month'], m['orders'], m['sales'], float(m['revenue'])])
    for cell in list(ws.columns)[3][1:]: cell.number_format = '#,##0.00'
    orders_sheet = wb.create_sheet('Orders', 0)
    orders_sheet.append(['Order / VIN', 'Customer', 'Package', 'Amount', 'Currency', 'Payment / order date', 'Status', 'Report', 'Seen', 'Customer email', 'Phone', 'Delivered at'])
    for order in sorted(data['orders'], key=lambda order: order['id']):
        orders_sheet.append([
            f"#{order['id']} / {order['vin']}",
            ' '.join(filter(None, [order.get('first_name'), order.get('last_name')])),
            (order.get('package') or '').title(), float(order.get('amount') or 0), order.get('currency') or 'USD',
            order.get('reporting_date'), (order.get('status') or '').title(),
            (order.get('report_status') or 'awaiting').title(), 'Seen' if order.get('admin_seen') else 'Unseen',
            order.get('email'), order.get('phone'), order.get('delivered_at')])
    for cell in list(orders_sheet.columns)[3][1:]: cell.number_format = '#,##0.00'
    wb.active = 0
    for title, records in [('Order Details', data['orders']), ('Customers', data['customers'])]:
        sheet = wb.create_sheet(title)
        # Preserve all database fields, including payment IDs, without exporting report contents.
        headers = list(records[0]) if records else (['id','vin','package','amount','status','customer_id','payment_id','created_at','paid_at','email','first_name','last_name','phone','reporting_date','has_report'] if title=='Order Details' else ['id','first_name','last_name','email','phone','created_at','order_count'])
        sheet.append(headers)
        for record in records:
            sheet.append([record.get(key) for key in headers])
        if 'amount' in headers:
            for cell in list(sheet.columns)[headers.index('amount')][1:]: cell.number_format = '#,##0.00'
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='112236')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Spreadsheet formula injection protection while preserving literal customer data.
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top')
                if cell.row % 2 == 0: cell.fill = PatternFill('solid', fgColor='EFF6FB')
        for i, column in enumerate(sheet.columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(55, max(14, max(len(str(c.value or '')) for c in column) + 2))
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def register_reporting(app, get_connection):
    @app.template_filter('admin_datetime')
    def admin_datetime(value):
        if not value:
            return '—'
        try:
            date = value if isinstance(value, datetime) else datetime.fromisoformat(value)
            return date.strftime('%b %d, %Y · %I:%M %p')
        except (TypeError, ValueError):
            return value

    def get_data():
        period = period_from_args(request.args)
        conn = get_connection()
        try:
            data = reporting_data(conn, period)
            data['fulfillment_alerts'] = [dict(r) for r in conn.execute(
                "SELECT a.*,o.vin FROM fulfillment_alerts a JOIN orders o ON o.id=a.order_id ORDER BY a.id DESC LIMIT 50")]
            data['fulfillment_pending'] = [dict(r) for r in conn.execute(
                "SELECT j.*,o.vin FROM fulfillment_jobs j JOIN orders o ON o.id=j.order_id WHERE j.state IN ('queued','retry','processing','sending','review','invalid_vin') ORDER BY j.order_id DESC LIMIT 50")]
            return data
        finally: conn.close()

    @app.route('/admin/dashboard')
    def admin_dashboard():
        if not session.get('admin_logged_in'): return redirect(url_for('admin_login'))
        response = app.make_response(render_template('admin_dashboard.html', **get_data()))
        response.headers['Cache-Control'] = 'no-store'
        return response

    @app.route('/admin/export-sales')
    def admin_export_sales():
        if not session.get('admin_logged_in'): return redirect(url_for('admin_login'))
        data = get_data()
        period = data['period']
        label = 'all-time' if period['months']=='all' else f"{period['month']}-{period['months']}months"
        response = send_file(build_workbook(data), as_attachment=True,
            download_name=f'VINLookupNow-sales-{label}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response.headers['Cache-Control'] = 'no-store'
        return response
